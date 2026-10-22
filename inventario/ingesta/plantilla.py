"""Estructura de la plantilla Excel que llenan las PYMES.

Cada hoja corresponde a una entidad del modelo de datos. Si se agrega
una columna aquí, el parser y la plantilla de ejemplo la toman solos.
"""
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HOJAS = {
    'Recepciones': [
        'fecha', 'nit_proveedor', 'proveedor', 'litros',
        'densidad', 'grasa_pct', 'temperatura',
    ],
    'Lotes': [
        'codigo_lote', 'producto', 'fecha_produccion',
        'fecha_vencimiento', 'cantidad',
    ],
    'Ventas': [
        'fecha', 'cliente', 'codigo_lote', 'producto',
        'cantidad', 'precio_unitario',
    ],
}


def _ejemplos():
    hoy = date.today()
    return {
        'Recepciones': [
            [hoy - timedelta(days=2), '900123456', 'Finca La Esperanza', 520, 1.031, 3.4, 4.0],
            [hoy - timedelta(days=1), '900654321', 'Finca El Roble', 410, 1.030, 3.6, 3.8],
        ],
        'Lotes': [
            ['L-001', 'Queso campesino', hoy - timedelta(days=2), hoy + timedelta(days=10), 80],
            ['L-002', 'Yogur natural 1 L', hoy - timedelta(days=1), hoy + timedelta(days=2), 120],
        ],
        'Ventas': [
            [hoy, 'Tienda La 14', 'L-001', 'Queso campesino', 20, 12000],
            [hoy, 'Supermercado Ubaté', 'L-002', 'Yogur natural 1 L', 35, 6500],
        ],
    }


def crear_plantilla(ruta, con_ejemplos=True):
    """Crea la plantilla Excel vacía o con filas de ejemplo."""
    libro = Workbook()
    libro.remove(libro.active)
    encabezado = Font(bold=True, color='FFFFFF')
    relleno = PatternFill('solid', fgColor='023A22')
    ejemplos = _ejemplos() if con_ejemplos else {}

    for nombre, columnas in HOJAS.items():
        hoja = libro.create_sheet(nombre)
        hoja.append(columnas)
        for celda in hoja[1]:
            celda.font = encabezado
            celda.fill = relleno
        for fila in ejemplos.get(nombre, []):
            hoja.append(fila)
        for i, columna in enumerate(columnas, start=1):
            hoja.column_dimensions[get_column_letter(i)].width = max(14, len(columna) + 4)
            if columna.startswith('fecha'):
                for celda in hoja[get_column_letter(i)][1:]:
                    celda.number_format = 'yyyy-mm-dd'
        hoja.freeze_panes = 'A2'

    libro.save(ruta)
    return ruta
